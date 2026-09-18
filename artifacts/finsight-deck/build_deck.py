from pathlib import Path
from colorsys import hls_to_rgb
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
OUT = Path(__file__).parent

def hsl(h,s,l):
    return ''.join(f'{round(v*255):02X}' for v in hls_to_rgb(h/360,l/100,s/100))
C = dict(bg=hsl(40,25,96), ink=hsl(24,14,12), card=hsl(40,22,98), muted=hsl(24,9,40), border=hsl(36,12,85), panel=hsl(24,10,11), paneltext=hsl(38,22,92), panelline=hsl(24,8,22), green=hsl(152,38,29), amber=hsl(36,74,40), red=hsl(5,62,44), navy='1E3A5F')
R = Presentation(); R.slide_width = Inches(13.333333); R.slide_height = Inches(7.5)
R.core_properties.title = 'FinSight — Early Warning to Intervention'
R.core_properties.subject = 'Northstar Commerce — fictional seeded financial demo'
R.core_properties.author = 'FinSight'
def rgb(c): return RGBColor.from_string(C.get(c,c))
def rect(s,x,y,w,h,fill='card',border=None):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=rgb(fill)
    from pptx.oxml.xmlchemy import OxmlElement
    a._element.spPr.append(OxmlElement('a:effectLst'))
    if border: a.line.color.rgb=rgb(border); a.line.width=Pt(.7)
    else: a.line.fill.background()
    return a

def text(s,content,x,y,w,h,size=20,color='ink',bold=False,font='Inter',align=PP_ALIGN.LEFT):
    a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=a.text_frame
    tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,part in enumerate(content.split('|')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=part
        p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=1.12
        for run in p.runs:
            run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=rgb(color)
    return a

def line(s,x1,y1,x2,y2,color='border',width=1):
    a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); a.line.color.rgb=rgb(color); a.line.width=Pt(width)
    from pptx.oxml.xmlchemy import OxmlElement
    a._element.spPr.append(OxmlElement('a:effectLst'))
    return a

def arrow(s,x1,y,x2,color='muted'):
    line(s,x1,y,x2,y,color,1.1); line(s,x2-.06,y-.05,x2,y,color,1.1); line(s,x2-.06,y+.05,x2,y,color,1.1)

def label(s,content,x,y,w,color='muted',size=10): return text(s,content.upper(),x,y,w,.22,size,color,True)

def slide(section,n,dark=False,foot='Northstar Commerce · fictional seeded demo'):
    s=R.slides.add_slide(R.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('panel' if dark else 'bg')
    fg='paneltext' if dark else 'ink'; dim='paneltext' if dark else 'muted'; ln='panelline' if dark else 'border'
    rect(s,.65,.42,.31,.31,fg); text(s,'F',.65,.439,.31,.25,13,'panel' if dark else 'bg',True,align=PP_ALIGN.CENTER)
    text(s,'FinSight',1.06,.417,1.42,.35,17,fg,True)
    line(s,2.5,.44,2.5,.71,ln); label(s,section,2.7,.475,8,dim,10)
    line(s,.65,.95,12.68,.95,ln); line(s,.65,6.96,12.68,6.96,ln)
    text(s,foot,.65,7.075,11,.19,9,dim)
    text(s,f'{n:02d} / 06',11.83,7.055,.85,.25,10,dim,font='JetBrains Mono',align=PP_ALIGN.RIGHT)
    return s

def notes(s,t): s.notes_slide.notes_text_frame.text=t

s=slide('Early warning → intervention',1)
text(s,'FinSight',.65,1.39,4.7,.78,49,bold=True)
text(s,'Autonomous Financial|Early-Warning &|Intervention System.',.68,2.42,4.62,1.54,25,bold=True)
rect(s,.68,4.29,.46,.045,'navy')
text(s,'What financial problem is forming,|why is it happening, and what|should we do before it|becomes expensive?',.68,4.63,4.64,1.48,18,'muted')
x,y,w,h=5.6,1.52,7.08,4.89
rect(s,x,y,w,h,'card','border'); rect(s,x,y,1.42,h,'panel')
text(s,'FinSight',x+.16,y+.23,1.18,.32,14,'paneltext',True)
text(s,'NORTHSTAR|COMMERCE',x+.16,y+.68,1.2,.5,8,'paneltext')
nav=['Overview','Risk Radar','Money Leaks','Cash Forecast','Simulator','Transactions','Vendors','Workflows','AI Analyst','Audit Trail']
for i,n in enumerate(nav):
    yy=y+1.38+i*.265
    if i==0: rect(s,x+.1,yy-.055,1.22,.255,'panelline')
    text(s,n,x+.18,yy,1.15,.22,8.1,'paneltext',i==0)
mx=x+1.64
text(s,'Overview',mx,y+.24,4.9,.36,19,bold=True)
text(s,'Northstar Commerce · INR',mx,y+.68,4.9,.26,10,'muted')
for xx,lab,val,tone in [(mx,'CURRENT CASH','₹4.82 Cr','ink'),(mx+2.61,'90-DAY FORECAST','₹2.17 Cr','amber')]:
    rect(s,xx,y+1.15,2.38,1.19,'card','border'); label(s,lab,xx+.16,y+1.31,2.1,size=8)
    text(s,val,xx+.16,y+1.69,2.1,.47,23,tone,True)
rect(s,mx,y+2.53,2.38,1.67,'panel')
label(s,'Financial health',mx+.17,y+2.72,2.05,'paneltext',8)
text(s,'78',mx+.17,y+3.07,1.05,.68,42,'paneltext',True)
text(s,'/100',mx+1.26,y+3.41,.89,.26,13,'paneltext',font='JetBrains Mono')
rect(s,mx+.17,y+3.94,2.03,.045,'panelline');rect(s,mx+.17,y+3.94,2.03*.78,.045,'green')
rect(s,mx+2.61,y+2.53,2.38,1.67,'card','border')
label(s,'Recoverable leakage',mx+2.77,y+2.72,2.1,size=8)
text(s,'₹38.4L',mx+2.77,y+3.17,2.08,.48,25,'green',True)
text(s,'Annualised',mx+2.77,y+3.82,2.0,.24,10,'muted')
text(s,'Selected Overview modules · editable UI mockup',5.61,6.56,7.08,.22,9,'muted')
notes(s,'FinSight asks what financial problem is forming, why it is happening, and what the team should do before it becomes expensive. Northstar Commerce is a fictional seeded demonstration. Overview combines current cash, a deterministic forecast, health and recoverable leakage. Production note: this is an editable reconstruction of selected existing UI modules, not a screenshot. The running preview was unreachable. Palette and typography were read from the application. At-risk capital was omitted because the brief and implementation disagree.')

s=slide('The problem',2)
text(s,'Finance teams see problems|after they become expensive.',.65,1.4,11.7,1.15,34,bold=True)
for yy,title,sub in [(3.06,'Reactive dashboards','Report the symptom after the fact.'),(4.08,'No causal explanation','An anomaly is not a diagnosis.'),(5.1,'No path to action','An alert is not an intervention.')]:
    rect(s,.68,yy+.08,.06,.26,'navy');text(s,title,.96,yy,5.45,.38,20,bold=True);text(s,sub,.96,yy+.44,5.45,.36,16,'muted')
rect(s,7.29,2.99,5.38,3.41,'panel')
label(s,'Northstar Commerce',7.6,3.28,4.75,'paneltext',10)
text(s,'52',7.53,3.69,2.86,1.42,98,'paneltext',True)
text(s,'days',10.42,4.48,1.84,.48,26,'paneltext')
text(s,'until a projected cash-threshold breach',7.61,5.34,4.65,.44,17,'paneltext')
text(s,'Receivables aging · safe floor ₹1.25 Cr',7.61,5.99,4.67,.24,11,'paneltext')
notes(s,'A dashboard can show a healthy-looking balance without explaining the risk forming underneath. Aging receivables reduce expected inflows. The seeded forecast projects a breach of the ₹1.25 Cr minimum-safe floor in 52 days. That is a warning horizon, not a realized loss. The day-90 balance can recover after an earlier trough; FinSight surfaces that interim breach rather than relying only on the endpoint.')

s=slide('How FinSight solves it',3)
text(s,'From a weak signal to a traceable action.',.65,1.42,12,.77,31,bold=True)
for i,st in enumerate(['Detect','Investigate','Explain','Predict','Simulate','Act']):
    xx=.65+i*2.055; rect(s,xx,2.55,1.75,.65,'panel' if i==5 else 'card','border' if i!=5 else None)
    text(s,st,xx,2.745,1.75,.32,16,'paneltext' if i==5 else 'ink',True,align=PP_ALIGN.CENTER)
    if i<5: arrow(s,xx+1.79,2.88,xx+2.00)
rect(s,.65,3.73,12.03,2.41,'card','border')
label(s,'Risk Radar / Trace cause',.93,3.98,8,'navy',11)
for i,c in enumerate(['Receivables|aging','Cash inflow|reduction','Cash reserve|decline','AP|exposure','Liquidity|risk']):
    xx=.93+i*2.35
    rect(s,xx,4.56,2.03,.91,'panel' if i==4 else 'bg','border' if i!=4 else None)
    text(s,c,xx+.12,4.765,1.79,.61,15,'paneltext' if i==4 else 'ink',True,align=PP_ALIGN.CENTER)
    if i<4: arrow(s,xx+2.08,5.025,xx+2.29)
text(s,'Trace the driver. Inspect the evidence. Choose the response.',.95,5.71,11.35,.27,13,'muted')
text(s,'Causal-chain UI excerpt · editable reconstruction, not a live screenshot',.65,6.42,12,.25,10,'muted')
notes(s,'FinSight connects detection to action. Follow receivables aging through reduced inflows, falling reserves, accounts-payable exposure and liquidity risk. The application presents evidence at the relevant nodes and lets the user create a workflow. This is a simplified reconstruction of Trace Cause; no transaction IDs, probabilities or confidence scores have been fabricated.')

s=slide('Two signature capabilities',4)
text(s,'Trace the cause. Test the response.',.65,1.42,12,.77,34,bold=True)
line(s,6.63,2.63,6.63,6.48)
label(s,'Financial Risk Graph',.67,2.61,5.5,'navy',11)
text(s,'Relationships,|not isolated rows.',.67,3.02,5.57,1.06,28,bold=True)
text(s,'Links transactions to downstream|financial consequences.',.67,4.27,5.5,.72,18,'muted')
for xx,yy,ww,tx in [(.68,5.48,1.54,'Receivables'),(2.64,5.48,1.5,'Cash'),(4.55,5.48,1.54,'Payables')]:
    rect(s,xx,yy,ww,.6,'card','border');text(s,tx,xx,yy+.18,ww,.3,13,bold=True,align=PP_ALIGN.CENTER)
arrow(s,2.27,5.78,2.57);arrow(s,4.18,5.78,4.49)
label(s,'Crisis Simulator',7.08,2.61,5.6,'navy',11)
text(s,'“What if revenue|drops 15%?”',7.08,3.02,5.57,1.06,28,bold=True)
text(s,'Compare before committing.',7.08,4.27,5.5,.4,18,'muted')
for yy,lab,col in [(5.0,'Baseline','ink'),(5.46,'Scenario','amber'),(5.92,'Scenario + recommended intervention','green')]:
    rect(s,7.1,yy+.03,.045,.27,col);text(s,lab,7.3,yy,5.33,.31,14,col,True)
text(s,'Deterministic comparison · hypothetical input, not a claimed result',7.08,6.48,5.59,.28,9,'muted')
notes(s,'The Financial Risk Graph connects transactions to financial consequences. The Crisis Simulator asks what changes if revenue drops 15 percent or another operating driver moves. Baseline, scenario and recommended-intervention forecasts are computed deterministically, with strategies ranked by the engine. Fifteen percent is a hypothetical input, not an observed result. No invented scenario balance, savings claim or decorative chart is shown.')

s=slide('Trust & architecture',5,True,foot='Northstar Commerce · fictional demo · EnterPro is a stateful mock workflow adapter')
label(s,'The safety principle',.65,1.43,11,'paneltext',11)
text(s,'Deterministic calculations are the source of truth.',.65,1.98,12.04,.69,28,'paneltext',True)
text(s,'Qwen explains; it does not calculate the numbers.',.65,2.81,12.02,.67,26,'paneltext')
for i,(name,sub) in enumerate([('Domain Engine','Calculations & forecasts'),('Qwen','Evidence-based explanation'),('EnterPro','Workflow → approval → audit')]):
    xx=.65+i*4.16
    rect(s,xx,4.18,3.69,1.51,'panel','panelline')
    text(s,name,xx+.25,4.47,3.19,.47,24,'paneltext',True)
    text(s,sub,xx+.25,5.17,3.19,.32,12,'paneltext')
    if i<2: arrow(s,xx+3.77,4.94,xx+4.07,'paneltext')
text(s,'Supplied metrics and evidence guide the narrative; financial outputs remain deterministic.',.65,6.22,12,.36,15,'paneltext')
notes(s,'The domain engine calculates; Qwen explains supplied metrics and evidence. Qwen is instructed not to invent numbers, which is not a guarantee of error-free prose. Deterministic financial outputs remain authoritative. EnterPro is currently a stateful mock adapter. Workflow, approval, leak-recovery and audit state are database-backed; this is not a connected external orchestration service or production authorization system.')

s=slide('Impact & close',6)
label(s,'Identify the opportunity. Surface the risk early.',.65,1.43,12,'navy',11)
text(s,'₹38.4L',.59,2.09,7.32,1.57,92,'green',True)
text(s,'recoverable annualised leakage identified',.68,3.78,7.15,.42,20,bold=True)
line(s,8.04,2.29,8.04,4.47)
text(s,'52',8.48,2.37,2.46,1.07,69,bold=True)
text(s,'days',10.9,2.92,1.56,.5,25)
text(s,'before a projected|cash-threshold breach',8.52,3.72,4.15,.7,18,'muted')
text(s,'AI-flagged risk → EnterPro workflow → approval → audit trail',.67,4.92,12,.41,18,'muted')
text(s,'FinSight doesn’t just report financial problems —|it finds them before they’re expensive, and helps the team act.',.67,5.69,12,1.01,24,bold=True)
notes(s,'FinSight identifies ₹38.4 lakh in annualised recoverable leakage and separately projects a cash-threshold breach in 52 days. Identified leakage is not realized savings; the warning is not proof of an averted breach. A flagged risk can proceed through the demonstrated approval and audit flow. FinSight does not just report financial problems: it finds them before they become expensive and helps the team act.')

# Keep native editable text using the app’s Inter and JetBrains Mono typefaces.
R.save(OUT/'FinSight-Pitch-Deck.pptx')
assert len(R.slides)==6
for i,s in enumerate(R.slides,1):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0, (i,sh.name,'negative coordinate')
        assert sh.left+sh.width<=R.slide_width+10 and sh.top+sh.height<=R.slide_height+10, (i,sh.name,'out of bounds')
        assert sh.shape_type != 13, (i,'unexpected raster image')
all_text=chr(10).join(sh.text for s in R.slides for sh in s.shapes if sh.has_text_frame)
assert '28.4' not in all_text and '28.6' not in all_text
(OUT/'slide-text.txt').write_text(all_text)
print('Created 6 editable 16:9 slides; all shapes within bounds; no rasterized text or charts.')
